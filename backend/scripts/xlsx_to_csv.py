# Converte a planilha de sazonalidade da CEAGESP (sazonalidade_quantidade-2023-2025.xlsx,
# obtida a partir da conversão do PDF de sazonalidade de alimentos publicado no site da
# CEAGESP) no CSV consolidado (sazonalidade_consolidada_produtos.csv) usado por seed.py.

# A planilha original marca a safra de cada mês pela COR de preenchimento da célula (sem
# texto), por isso o script lê o HEX da célula e o traduz pra uma categoria de safra.

import os
import openpyxl
from openpyxl.styles import PatternFill
import pandas as pd

# CONFIGURAÇÕES E MAPEAMENTOS
DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'data')
ARQUIVO_ENTRADA = os.path.join(DATA_DIR, 'sazonalidade_quantidade-2023-2025.xlsx')
ARQUIVO_SAIDA = os.path.join(DATA_DIR, 'sazonalidade_consolidada_produtos.csv')
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']

# Escala numérica para cálculo de média
MAPA_NUMERICO = {
    "Safra Alta": 4,
    "Safra Média/Alta": 3,
    "Safra Baixa/Média": 2,
    "Safra Muito Baixa": 1,
    "Sem Volume": 0
}
MAPA_TEXTO = {v: k for k, v in MAPA_NUMERICO.items()}

# FUNÇÕES DE APOIO
def obter_hex_da_celula(celula):
    """Extrai o código Hexadecimal da célula, ignorando brancos/nulos."""
    fill = celula.fill
    if not hasattr(fill, 'start_color'): return ""
    idx = str(fill.start_color.index)
    if idx in ['00000000', '0', 'FFFFFFFF', 'system_window_text']:
        return ""
    return f"#{idx[-6:]}".upper()

def categorizar_safra(hex_cor):
    """Transforma códigos HEX em categorias baseadas na temperatura da cor."""
    if not hex_cor: return "Sem Volume"
    # Tons Quentes (Vermelhos/Corais)
    if any(hex_cor.startswith(p) for p in ['#F3', '#F4', '#F5']):
        return "Safra Muito Baixa"
    # Tons Médios (Laranjas/Amarelos)
    elif hex_cor.startswith('#F'): 
        return "Safra Baixa/Média"
    # Tons Transição (Creme/Verde Lima)
    elif hex_cor.startswith('#D') or hex_cor.startswith('#C'):
        return "Safra Média/Alta"
    # Tons Frios (Verdes Fortes)
    else: 
        return "Safra Alta"

def mapear_cores_unicas(sheet, arquivo_saida='cores_unicas.xlsx'):
    """Varre todas as células da planilha e salva um .xlsx com cada cor HEX
    encontrada, sua contagem, um exemplo de célula e um preenchimento com
    a própria cor ao lado, para conferência visual."""
    cores = {}
    for row in sheet.iter_rows():
        for celula in row:
            hex_cor = obter_hex_da_celula(celula)
            if not hex_cor:
                continue
            if hex_cor not in cores:
                cores[hex_cor] = {"contagem": 0, "exemplo_celula": celula.coordinate}
            cores[hex_cor]["contagem"] += 1

    wb_saida = openpyxl.Workbook()
    ws_saida = wb_saida.active
    ws_saida.title = "Cores Únicas"
    ws_saida.append(["hex", "cor", "contagem", "exemplo_celula"])

    for hex_cor, info in sorted(cores.items()):
        ws_saida.append([hex_cor, "", info["contagem"], info["exemplo_celula"]])
        celula_cor = ws_saida.cell(row=ws_saida.max_row, column=2)
        celula_cor.fill = PatternFill(start_color=hex_cor[1:], end_color=hex_cor[1:], fill_type="solid")

    ws_saida.column_dimensions['A'].width = 12
    ws_saida.column_dimensions['B'].width = 8
    ws_saida.column_dimensions['D'].width = 15

    wb_saida.save(arquivo_saida)
    print(f"- Encontradas {len(cores)} cores únicas. Salvo em '{arquivo_saida}'.")
    return cores

# PASSO 1: LEITURA E EXTRAÇÃO (POR VARIEDADE)

wb = openpyxl.load_workbook(ARQUIVO_ENTRADA, data_only=True)
sheet = wb.active

# mapear_cores_unicas(sheet)

dados_brutos = []
for row in sheet.iter_rows(min_row=1): 

    val_produto = str(row[2].value).strip().upper() if row[2].value else ""
    val_grupo = str(row[0].value).strip().upper() if row[0].value else ""
    
    if len(row) < 17 or val_produto in ["", "PRODUTO"]:
        continue
        
    # Ignora se for grupo de flores
    if "FLORES" in val_grupo:
        continue
    
    item = {
        "Grupo": row[0].value,
        "Produto": row[2].value,
        "Variedade": row[4].value
    }
    
    meses_sem_volume = 0
    for i, mes in enumerate(MESES):
        cor_hex = obter_hex_da_celula(row[i + 5])
        categoria = categorizar_safra(cor_hex)
        item[mes] = categoria
        if categoria == "Sem Volume":
            meses_sem_volume += 1
            
    if meses_sem_volume > 3:
        continue
        
    dados_brutos.append(item)

df_variedades = pd.DataFrame(dados_brutos)


# PASSO 2: AGRUPAMENTO E MÉDIA (POR PRODUTO)

# 1. Converte categorias para números para poder calcular
df_calc = df_variedades.copy()
for mes in MESES:
    df_calc[mes] = df_calc[mes].map(MAPA_NUMERICO)

# 2. Agrupa por Produto e Grupo, tirando a média das variedades
# Usamos round() para voltar ao nível de categoria mais próximo
df_agrupado = df_calc.groupby(['Grupo', 'Produto'])[MESES].mean().round().astype(int)

# 3. Converte números de volta para texto para o relatório final
df_final = df_agrupado.copy()
for mes in MESES:
    df_final[mes] = df_final[mes].map(MAPA_TEXTO)

df_final = df_final.reset_index()

# PASSO 3: EXPORTAÇÃO
# df_variedades.to_csv(os.path.join(DATA_DIR, 'sazonalidade_detalhada_variedades.csv'), index=False, encoding='utf-8-sig')
df_final.to_csv(ARQUIVO_SAIDA, index=False, encoding='utf-8-sig')

print(f"Sucesso!")
print(f"- Processados {len(df_variedades)} itens detalhados.")
print(f"- Gerados {len(df_final)} produtos agrupados.")
print("\nPrimeiros resultados (Agrupados):")
print(df_final.head())